import sys
import os
import argparse
import json
import random
import math
from datetime import datetime
from PIL import Image, ImageDraw, ImageFont

def get_font(font_path, size):
    try:
        return ImageFont.truetype(font_path, size)
    except Exception:
        try:
            return ImageFont.load_default()
        except Exception:
            return None

def wrap_text_to_width(text, font, max_width, draw):
    words = text.split(' ')
    lines = []
    current_line = []
    
    for word in words:
        test_line = ' '.join(current_line + [word])
        bbox = draw.textbbox((0, 0), test_line, font=font)
        line_w = bbox[2] - bbox[0]
        if line_w <= max_width:
            current_line.append(word)
        else:
            if current_line:
                lines.append(' '.join(current_line))
                current_line = [word]
            else:
                lines.append(word)
                current_line = []
    if current_line:
        lines.append(' '.join(current_line))
    return lines

def draw_hand_text_line(draw, text, start_x, start_y, font, base_rgb, rng):
    words = text.split(' ')
    curr_x = start_x
    for i, word in enumerate(words):
        if not word:
            continue
        # Baseline micro-jitter: replicates organic human hand tremor (-1.2 to +1.2 px)
        jitter_y = start_y + rng.uniform(-1.2, 1.2)
        
        # Micro-variation in ink density: simulates varying ballpoint pen pressure
        r = max(0, min(255, int(base_rgb[0] + rng.randint(-8, 8))))
        g = max(0, min(255, int(base_rgb[1] + rng.randint(-8, 8))))
        b = max(0, min(255, int(base_rgb[2] + rng.randint(-8, 8))))
        ink_color = (r, g, b)

        draw.text((curr_x, jitter_y), word, fill=ink_color, font=font)
        bbox = draw.textbbox((curr_x, jitter_y), word, font=font)
        word_w = bbox[2] - bbox[0]
        
        # Organic spacing variance between words
        space_w = font.size * 0.30 + rng.uniform(-0.6, 1.2)
        curr_x += word_w + space_w

def draw_hand_underline(draw, x1, y, x2, color, rng, width=2):
    # Organic wavy underline beneath headings and Ans tags
    steps = max(3, int((x2 - x1) / 18))
    step_w = (x2 - x1) / steps
    points = []
    for i in range(steps + 1):
        px = x1 + i * step_w
        py = y + rng.uniform(-1.0, 1.0)
        points.append((px, py))
    for i in range(len(points) - 1):
        draw.line([points[i], points[i+1]], fill=color, width=width)

def draw_binder_holes(draw, height, hole_x=65, radius=16):
    # 3 Realistic punched binder holes along the notebook margin
    ys = [int(height * 0.18), int(height * 0.50), int(height * 0.82)]
    for y in ys:
        draw.ellipse([(hole_x - radius - 2, y - radius - 2), (hole_x + radius + 2, y + radius + 2)], fill='#ded7cb')
        draw.ellipse([(hole_x - radius, y - radius), (hole_x + radius, y + radius)], fill='#f5efe6', outline='#c8beaf', width=2)

def generate_handwritten_pages(assignment_title, subject_name, student_name, reg_no, qa_items, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    width, height = 1240, 1754 # A4 at 150 DPI
    line_height = 46
    top_margin = 240
    left_margin = 185
    right_margin = width - 65
    content_width = right_margin - (left_margin + 20)
    bottom_margin = height - 90

    # Fonts
    font_bold = get_font('C:/Windows/Fonts/segoeprb.ttf', 24) or get_font('C:/Windows/Fonts/arialbd.ttf', 22)
    font_header = get_font('C:/Windows/Fonts/segoepr.ttf', 22) or get_font('C:/Windows/Fonts/arial.ttf', 20)
    font_body = get_font('C:/Windows/Fonts/Inkfree.ttf', 28) or get_font('C:/Windows/Fonts/segoepr.ttf', 24)
    font_question = get_font('C:/Windows/Fonts/segoeprb.ttf', 25) or get_font('C:/Windows/Fonts/arialbd.ttf', 22)

    # Inks
    ink_black_rgb = (20, 27, 40)       # Black Gel Pen (Questions & Headers)
    ink_burgundy_rgb = (136, 19, 55)   # Dark Crimson / Burgundy (Ans tags & underlines)
    ink_blue_rgb = (22, 52, 102)       # Royal Blue Ballpoint Pen (Answers)

    # Temporary image for measuring
    meas_img = Image.new('RGB', (width, height))
    meas_draw = ImageDraw.Draw(meas_img)

    # Flatten all content into discrete printable lines
    printable_lines = [] # list of (type, tag, text)
    for idx, item in enumerate(qa_items, start=1):
        q_raw = item.get('q', f'Question {idx}')
        a_raw = item.get('a', '')
        
        # Wrap question
        q_tag = f"Q{idx}."
        q_lines = wrap_text_to_width(q_raw, font_question, content_width - 10, meas_draw)
        if q_lines:
            printable_lines.append(('q_head', q_tag, q_lines[0]))
            for ql in q_lines[1:]:
                printable_lines.append(('q_cont', '', ql))
        else:
            printable_lines.append(('q_head', q_tag, ''))
        
        printable_lines.append(('ans_head', 'Ans:', ''))
        
        # Split answers by paragraph then wrap
        paragraphs = a_raw.split('\n')
        for p in paragraphs:
            p = p.strip()
            if not p:
                printable_lines.append(('blank', '', ''))
                continue
            wrapped_p = wrap_text_to_width(p, font_body, content_width, meas_draw)
            for wl in wrapped_p:
                printable_lines.append(('a', '', wl))
        
        printable_lines.append(('blank', '', ''))
        printable_lines.append(('blank', '', ''))

    # Calculate pages needed
    max_lines_per_page = (bottom_margin - top_margin) // line_height - 1
    pages_data = []
    curr_page = []
    
    for item in printable_lines:
        if len(curr_page) >= max_lines_per_page:
            pages_data.append(curr_page)
            curr_page = []
        curr_page.append(item)
    if curr_page:
        pages_data.append(curr_page)

    if not pages_data:
        pages_data = [[('ans_head', 'Ans:', 'No content specified.')]]

    total_pages = len(pages_data)
    generated_images = []
    today_str = datetime.now().strftime("%d/%m/%Y")

    for page_idx, page_lines in enumerate(pages_data, start=1):
        rng = random.Random(page_idx * 1337 + 42)
        img = Image.new('RGB', (width, height), color='#fdfbf6')
        draw = ImageDraw.Draw(img)

        # Draw 3 binder holes on left margin
        draw_binder_holes(draw, height, hole_x=65, radius=16)

        # Draw notebook horizontal ruled lines
        for y in range(top_margin, bottom_margin + line_height, line_height):
            draw.line([(left_margin - 60, y), (right_margin, y)], fill='#c4d7ec', width=2)

        # Draw red double left margin
        draw.line([(left_margin, 70), (left_margin, height - 45)], fill='#e05d5d', width=3)
        draw.line([(left_margin - 8, 70), (left_margin - 8, height - 45)], fill='#f5bcbc', width=1)

        # Top Header Box
        draw_hand_text_line(draw, f"NAME: {student_name.upper()}", left_margin + 20, 95, font_bold, ink_black_rgb, rng)
        draw_hand_text_line(draw, f"REG NO: {reg_no.upper()}", left_margin + 20, 138, font_bold, ink_black_rgb, rng)
        draw_hand_text_line(draw, f"SUBJECT: {subject_name.upper()}", width - 450, 95, font_header, ink_black_rgb, rng)
        draw_hand_text_line(draw, f"PAGE: {page_idx:02d} / {total_pages:02d}    DATE: {today_str}", width - 450, 138, font_header, ink_black_rgb, rng)
        
        # Organic hand-drawn header separator rule with ample vertical breathing room
        draw_hand_underline(draw, left_margin - 30, 195, right_margin, '#64748b', rng, width=2)

        # Draw content
        curr_y = top_margin + 6

        for l_type, l_tag, l_text in page_lines:
            if l_type == 'q_head':
                # Question tag in margin column (e.g. Q1.)
                draw_hand_text_line(draw, l_tag, left_margin - 60, curr_y, font_question, ink_black_rgb, rng)
                # Question text in main ruled notebook column
                draw_hand_text_line(draw, l_text, left_margin + 20, curr_y, font_question, ink_black_rgb, rng)
            elif l_type == 'q_cont':
                # Continuation line of question
                draw_hand_text_line(draw, l_text, left_margin + 20, curr_y, font_question, ink_black_rgb, rng)
            elif l_type == 'ans_head':
                # Burgundy Ans tag in main column with organic hand-drawn wavy underline
                draw_hand_text_line(draw, l_tag, left_margin + 20, curr_y, font_bold, ink_burgundy_rgb, rng)
                ans_w = draw.textbbox((0, 0), l_tag, font=font_bold)[2]
                draw_hand_underline(draw, left_margin + 20, curr_y + 36, left_margin + 20 + ans_w + 12, '#881337', rng, width=2)
            elif l_type == 'a':
                # Royal Blue Ballpoint Pen with human baseline tremor and variable ink pressure
                draw_hand_text_line(draw, l_text, left_margin + 20, curr_y, font_body, ink_blue_rgb, rng)
            
            curr_y += line_height

        page_path = os.path.join(output_dir, f"page_{page_idx:02d}.png")
        img.save(page_path, dpi=(150, 150))
        generated_images.append(page_path)

    # Assemble into PDF
    clean_title = "".join(c for c in assignment_title if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    clean_subj = "".join(c for c in subject_name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    clean_reg = "".join(c for c in reg_no if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    
    pdf_filename = f"{clean_title}_{clean_subj}_{clean_reg}.pdf"
    pdf_path = os.path.join(output_dir, pdf_filename)

    if generated_images:
        pil_images = [Image.open(p).convert('RGB') for p in generated_images]
        pil_images[0].save(pdf_path, save_all=True, append_images=pil_images[1:], resolution=150)
        print(f"SUCCESS_PDF: {pdf_path}")

    return {
        "pdf_path": pdf_path,
        "images": generated_images,
        "total_pages": total_pages,
        "pdf_filename": pdf_filename
    }

def generate_word_document(assignment_title, subject_name, student_name, reg_no, qa_items, output_dir):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    os.makedirs(output_dir, exist_ok=True)
    doc = docx.Document()

    # Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(assignment_title.upper())
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Student metadata table
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Light Shading'
    table.rows[0].cells[0].text = f"Student Name: {student_name}"
    table.rows[0].cells[1].text = f"Registration No: {reg_no}"
    table.rows[1].cells[0].text = f"Subject: {subject_name}"
    table.rows[1].cells[1].text = "Format: Standard College Assignment"

    doc.add_paragraph() # Spacing

    for idx, item in enumerate(qa_items, start=1):
        q_raw = item.get('q', f'Question {idx}')
        a_raw = item.get('a', '')

        qp = doc.add_paragraph()
        q_run = qp.add_run(f"Q{idx}. {q_raw}")
        q_run.bold = True
        q_run.font.size = Pt(13)
        q_run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

        ap = doc.add_paragraph()
        ans_badge = ap.add_run("Answer:\n")
        ans_badge.bold = True
        ans_badge.font.size = Pt(11)
        ans_badge.font.color.rgb = RGBColor(0x83, 0x18, 0x43)

        body_run = ap.add_run(a_raw)
        body_run.font.size = Pt(11)
        body_run.font.color.rgb = RGBColor(0x16, 0x31, 0x5c)

        doc.add_paragraph() # Separator

    clean_title = "".join(c for c in assignment_title if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    clean_subj = "".join(c for c in subject_name if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    clean_reg = "".join(c for c in reg_no if c.isalnum() or c in (' ', '_', '-')).strip().replace(' ', '_')
    docx_filename = f"{clean_title}_{clean_subj}_{clean_reg}.docx"
    docx_path = os.path.join(output_dir, docx_filename)

    doc.save(docx_path)
    print(f"SUCCESS_DOCX: {docx_path}")

    return {
        "docx_path": docx_path,
        "filename": docx_filename,
        "format": "word"
    }

def main():
    parser = argparse.ArgumentParser(description="NEXORA Assignment Generator Engine")
    parser.add_argument("--mode", default="handwritten", choices=["handwritten", "word"])
    parser.add_argument("--name", default="Student")
    parser.add_argument("--regno", default="REG-0001")
    parser.add_argument("--subject", default="Computer Science")
    parser.add_argument("--assignment", default="Assignment_1")
    parser.add_argument("--data-json", help="Path to JSON file containing qa items or raw json string")
    parser.add_argument("--output-dir", default=os.path.join(os.path.expanduser("~"), "Desktop"))

    args = parser.parse_args()

    qa_items = []
    if args.data_json:
        if os.path.exists(args.data_json):
            with open(args.data_json, 'r', encoding='utf-8') as f:
                qa_items = json.load(f)
        else:
            try:
                qa_items = json.loads(args.data_json)
            except Exception:
                pass

    if not qa_items:
        # Default starter QA if none supplied
        qa_items = [
            {
                "q": "Explain the difference between Process and Thread in modern Operating Systems.",
                "a": "1. A process is an executing program instance with its own virtual address space, memory pages, and system handles.\n2. A thread is an execution unit inside a process that shares the parent process heap, global data, and code space.\n3. Inter-thread communication is faster and more resource-efficient than inter-process communication (IPC).\n4. A process crash usually isolates system damage, whereas an unhandled exception in one thread can terminate the entire process."
            },
            {
                "q": "What is Virtual Memory and why is Demand Paging critical?",
                "a": "1. Virtual memory is an abstraction provided by the OS memory manager and MMU that gives each process the illusion of large, contiguous physical memory.\n2. Demand paging only loads memory pages into physical RAM frames when referenced by the CPU.\n3. If a referenced page is not in RAM, a page fault occurs, loading the requested page from swap storage into RAM."
            }
        ]

    if args.mode == "word":
        result = generate_word_document(
            assignment_title=args.assignment,
            subject_name=args.subject,
            student_name=args.name,
            reg_no=args.regno,
            qa_items=qa_items,
            output_dir=args.output_dir
        )
    else:
        result = generate_handwritten_pages(
            assignment_title=args.assignment,
            subject_name=args.subject,
            student_name=args.name,
            reg_no=args.regno,
            qa_items=qa_items,
            output_dir=args.output_dir
        )

    print(json.dumps(result))

if __name__ == "__main__":
    main()

